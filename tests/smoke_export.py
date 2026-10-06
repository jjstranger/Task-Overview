"""财务导出的「范围 / 项目 / 内容 / 文件名」回归（db 层，不开窗口）。

对应界面上的 #ovExport 弹窗与 api.finance_export_save。要点：

1. `norm_opts` 的默认值 = 老行为（全部时间 / 全部项目 / 四项内容全要），
   所以老的 `export(fmt, year)` 调用方不用改；
2. 时间范围是**闭区间**，且会把「没填日期」的流水排除掉 —— 这种流水归不到
   任何一段时间里，塞进来会让人以为"这段时间真有这笔"；
3. `project_ids` / `client_id` 各自独立，能叠；
4. 内容勾选真的影响产出：xlsx 少表、json 少键；CSV 恒只有明细（后端兜底，
   否则界面上勾了项目汇总却导不出来）；
5. 四项全不勾要**报错**，不能悄悄退回"全要"；
6. 文件名：范围 tag + 单项目名（非法字符要洗掉，Windows 上 `/` `:` 会让
   系统另存为对话框直接报路径非法）+ 扩展名。

跑法：
    <venv>\\Scripts\\python.exe tests\\smoke_export.py
"""
from __future__ import annotations

import csv
import json
import os
import sys
import tempfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "app"))

from db import Db            # noqa: E402
from finance import Finance  # noqa: E402
from models import Board     # noqa: E402

FAILS = []
N = [0]


def chk(name, cond):
    N[0] += 1
    print(("  PASS  " if cond else "  FAIL  ") + name)
    if not cond:
        FAILS.append(name)


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="board_export_")
    out = os.path.join(tmp, "out")
    os.makedirs(out, exist_ok=True)
    db = Db(os.path.join(tmp, "board.sqlite"))
    db.open()
    board = Board(db)
    fin = Finance(db)

    c1 = fin.client_save(None, {"name": "甲客户"})
    c2 = fin.client_save(None, {"name": "乙客户"})
    # 项目名里故意带 `/`：文件名要洗掉它
    p1 = board.add_project("项目/甲", "commercial", c1)
    p2 = board.add_project("项目乙", "commercial", c2)

    fin.add({"kind": "contract", "project_id": p1, "amount": 100000,
             "date": "2025-03-01", "status": "有效"})
    fin.add({"kind": "payment", "project_id": p1, "amount": 30000,
             "date": "2026-01-15", "status": "已收"})
    fin.add({"kind": "payment", "project_id": p2, "amount": 20000,
             "date": "2026-06-30", "status": "已收"})
    fin.add({"kind": "outsource", "project_id": p2, "amount": 5000,
             "date": "2026-06-30", "status": "已付"})
    # 没填日期的那笔：改老库/导入数据里真实存在，只能直接写库造出来
    # （fin.add 会把空日期补成今天）
    db.run("INSERT INTO finance(project_id,node_id,kind,amount,currency,date,status)"
           " VALUES(?,?,?,?,?,?,?)", (p1, None, "payment", 9999, "CNY", "", "已收"))

    # ---- 1. norm_opts 默认值 = 老行为 ----
    o = fin.norm_opts(None)
    chk("默认 opts：xlsx + 四项内容全要 + 不限时间/项目",
        o["fmt"] == "xlsx" and o["sections"] == ["records", "projects", "clients", "aging"]
        and o["year"] == 0 and o["date_from"] == "" and o["date_to"] == ""
        and o["project_ids"] == [] and o["client_id"] == 0)
    chk("认不出来的格式退回 xlsx",
        fin.norm_opts({"fmt": "txt"})["fmt"] == "xlsx"
        and fin.norm_opts({"fmt": ""})["fmt"] == "xlsx")
    chk("CSV 强制只导款项明细（勾了别的也白勾）",
        fin.norm_opts({"fmt": "csv", "sections": ["projects", "aging"]})["sections"]
        == ["records"])
    try:
        fin.norm_opts({"sections": []})
        chk("四项全不勾要报错（不能悄悄退回全要）", False)
    except ValueError:
        chk("四项全不勾要报错（不能悄悄退回全要）", True)

    # ---- 2. 范围文案 ----
    chk("范围文案：不设 = 全部时间", fin.range_label(fin.norm_opts({})) == "全部时间")
    chk("范围文案：年份 / 起 / 止 / 闭区间",
        fin.range_label(fin.norm_opts({"year": 2026})) == "2026 年"
        and fin.range_label(fin.norm_opts({"date_from": "2026-01-01"})) == "2026-01-01 起"
        and fin.range_label(fin.norm_opts({"date_to": "2026-06-30"})) == "截至 2026-06-30"
        and fin.range_label(fin.norm_opts({"date_from": "2026-01-01",
                                           "date_to": "2026-06-30"}))
        == "2026-01-01 至 2026-06-30")

    # ---- 3. 过滤真的生效 ----
    def rows(**kw):
        return fin.pick_rows(fin.norm_opts(kw))

    chk("不筛 = 5 笔（含没日期那笔）", len(rows()) == 5)
    rng = rows(date_from="2026-01-01", date_to="2026-06-30")
    chk("时间范围是闭区间、且排除没日期的流水（实得 %d 笔）" % len(rng),
        len(rng) == 3
        and {r["date"] for r in rng} == {"2026-01-15", "2026-06-30"}
        and all(r["date"] for r in rng))
    chk("只给起始日期也生效（2025-03-01 起 = 4 笔，没日期那笔仍在外面）",
        len(rows(date_from="2025-03-01")) == 4)
    chk("年份筛选仍可用（2025 年 = 1 笔）", len(rows(year=2025)) == 1)
    chk("指定单个项目（项目乙 = 2 笔）",
        len(rows(project_ids=[p2])) == 2
        and all(r["project_id"] == p2 for r in rows(project_ids=[p2])))
    chk("指定多个项目 = 甲乙全部 5 笔（甲 3 + 乙 2，没日期那笔也算）",
        len(rows(project_ids=[p1, p2])) == 5
        and len(rows(project_ids=[p1])) == 3)
    chk("客户与项目两个条件各自独立、能叠",
        len(rows(client_id=c1)) == 3
        and len(rows(client_id=c1, project_ids=[p1])) == 3
        and len(rows(client_id=c1, project_ids=[p2])) == 0)

    # ---- 4. 文件名 ----
    n_all = fin.build_name("xlsx", fin.norm_opts({}))
    chk("默认文件名 = board_finance_all_<时间戳>.xlsx（实得 %s）" % n_all,
        n_all.startswith("board_finance_all_") and n_all.endswith(".xlsx"))
    chk("文件名带年份 tag",
        "_2026_" in fin.build_name("xlsx", fin.norm_opts({"year": 2026})))
    chk("文件名带闭区间 tag",
        "_20260101-20260630_"
        in fin.build_name("xlsx", fin.norm_opts({"date_from": "2026-01-01",
                                                 "date_to": "2026-06-30"})))
    n_one = fin.build_name("xlsx", fin.norm_opts({"project_ids": [p1]}))
    chk("单项目导出把项目名写进文件名、非法字符被洗掉（实得 %s）" % n_one,
        "_项目_甲_" in n_one and "/" not in n_one)
    chk("多项目导出只写个数（实得 %s）"
        % fin.build_name("csv", fin.norm_opts({"project_ids": [p1, p2]})),
        "_2prj_" in fin.build_name("csv", fin.norm_opts({"project_ids": [p1, p2]}))
        and fin.build_name("csv", fin.norm_opts({"project_ids": [p1, p2]})).endswith(".csv"))

    # ---- 5. 内容勾选真的影响产出 ----
    from openpyxl import load_workbook

    def xlsx_sheets(**kw):
        p = os.path.join(out, "a.xlsx")
        fin.export_to(p, "xlsx", kw)
        return load_workbook(p).sheetnames

    chk("只勾项目汇总 → 只有一张表",
        xlsx_sheets(sections=["projects"]) == ["项目汇总"])
    chk("勾明细 + 账龄 → 两张表、顺序跟界面一致",
        xlsx_sheets(sections=["records", "aging"]) == ["款项明细", "应收账龄"])
    chk("默认（不传 sections）→ 四张表，跟老行为一致",
        xlsx_sheets() == ["款项明细", "项目汇总", "客户汇总", "应收账龄"])

    pj = os.path.join(out, "b.xlsx")
    fin.export_to(pj, "xlsx", {"sections": ["projects"], "project_ids": [p2]})
    ws = load_workbook(pj)["项目汇总"]
    chk("内容勾选 + 项目筛选一起生效（项目汇总只有 1 行数据）", ws.max_row == 2)

    jp = os.path.join(out, "a.json")
    fin.export_to(jp, "json", {"sections": ["projects"]})
    with open(jp, encoding="utf-8") as f:
        j = json.load(f)
    chk("json 只勾项目汇总 → 只有 projects 键（+ 头信息）",
        set(j) == {"exported_at", "range", "year", "projects"})

    fin.export_to(jp, "json", {"sections": ["records"]})
    with open(jp, encoding="utf-8") as f:
        j = json.load(f)
    chk("json 勾明细 → 带 cards 和 records（cards 就是这批数的汇总口径）",
        set(j) == {"exported_at", "range", "year", "cards", "records"} and j["year"] == 0)
    chk("json 里的 range 是人话", j["range"] == "全部时间")

    cp = os.path.join(out, "a.csv")
    fin.export_to(cp, "csv", {"sections": ["projects"],
                              "date_from": "2026-01-01", "date_to": "2026-06-30"})
    with open(cp, encoding="utf-8-sig", newline="") as f:
        body = list(csv.reader(f))
    chk("CSV 是一张平表：11 列（勾了项目汇总也只给明细）", len(body[0]) == 11)
    chk("CSV 跟着时间范围走（1 行表头 + 3 行数据，实得 %d）" % (len(body) - 1),
        len(body) == 4)

    # ---- 6. 老路径不受影响 ----
    legacy = fin.export("xlsx")
    chk("老的 export(fmt) 仍写进导出目录",
        os.path.dirname(os.path.abspath(legacy)) == os.path.abspath(fin.export_dir()))
    chk("老路径出来的还是四张表",
        load_workbook(legacy).sheetnames
        == ["款项明细", "项目汇总", "客户汇总", "应收账龄"])

    db.close()
    print()
    if FAILS:
        print("FAIL %d / %d" % (len(FAILS), N[0]))
        for x in FAILS:
            print("   -", x)
        return 1
    print("ALL PASS (%d)" % N[0])
    return 0


if __name__ == "__main__":
    sys.exit(main())
