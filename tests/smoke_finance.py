"""M3 财务回归测试：口径、账龄、环节挂载、三种导出格式。

跑法：
    <venv>\\Scripts\\python.exe tests\\smoke_finance.py
"""
from __future__ import annotations

import os
import sys
import tempfile
from datetime import date, timedelta

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "app"))

from db import Db  # noqa: E402
from finance import Finance, money  # noqa: E402
from models import Board  # noqa: E402


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="board_fin_")
    path = os.path.join(tmp, "board.sqlite")
    db = Db(path)
    db.open()
    board = Board(db)
    fin = Finance(db)

    # ---- 准备数据 ----
    cid = fin.client_save(None, {"name": "ACME 影业", "contact": "王工",
                                 "settlement_cycle": "月结 30"})
    # 外包商（2026-10-16 打通）：外包支出是**必填**这个人的，所以先建两条
    v_a = fin.vendor_save(None, {"name": "阿强", "kind": "个人", "contact": "微信 qiang"})
    v_b = fin.vendor_save(None, {"name": "光影工作室", "kind": "团队"})
    pid = board.add_project("测试项目 A", "commercial", cid)
    stage = board.add_node(pid, None, "第一集")
    sub = board.add_node(pid, stage, "合成")

    today = date.today()
    fin.add({"kind": "contract", "project_id": pid, "amount": 100000,
             "date": (today - timedelta(days=45)).isoformat(), "status": "有效"})
    fin.add({"kind": "invoice", "project_id": pid, "amount": 60000,
             "date": (today - timedelta(days=20)).isoformat(), "status": "已开",
             "invoice_no": "INV-001"})
    fin.add({"kind": "payment", "project_id": pid, "amount": 40000,
             "date": (today - timedelta(days=10)).isoformat(), "status": "已收"})
    # 挂在二级子任务上的到账（分批款）——验证深层节点也能挂
    fin.add({"kind": "payment", "project_id": pid, "node_id": sub, "amount": 10000,
             "date": (today - timedelta(days=3)).isoformat(), "status": "已收"})
    # 待收，日期靠后
    fin.add({"kind": "payment", "project_id": pid, "amount": 50000,
             "date": (today + timedelta(days=15)).isoformat(), "status": "待收"})
    # 未开票部分应能被算出来：合同 100000 - 已开 60000 = 40000

    # ---- 口径 ----
    d = fin.data()
    c = d["cards"]
    assert c["contract"] == 100000, c
    assert c["received"] == 50000, c            # 40000 + 10000（含环节那笔）
    assert c["pending"] == 50000, c             # 100000 - 50000
    assert c["invoiced"] == 60000, c
    assert c["uninvoiced"] == 40000, c
    assert c["rate"] == 50.0, c

    # ---- 外包支出：付给外协的钱要从收入里扣掉，但只有「已付」才扣 ----
    fin.add({"kind": "outsource", "project_id": pid, "amount": 30000,
             "date": today.isoformat(), "status": "已付", "note": "给外协",
             "vendor_id": v_a})
    fin.add({"kind": "outsource", "project_id": pid, "amount": 7000,
             "date": today.isoformat(), "status": "应付", "note": "还没付",
             "vendor_id": v_b})
    d2 = fin.data()
    c2 = d2["cards"]
    assert c2["outsource"] == 30000, c2               # 应付那笔不算
    assert c2["net"] == 20000, c2                     # 已收 50000 − 外包 30000
    assert c2["received"] == 50000, c2                # 外包不该动「已收」
    assert c2["pending"] == 50000, c2                 # 也不该动「待收」
    pj0 = d2["projects"][0]
    assert pj0["outsource"] == 30000 and pj0["net"] == 20000, pj0
    assert d2["client_summary"][0]["net"] == 20000, d2["client_summary"]
    # 改成「应付」→ 不再冲减；改回「已付」→ 又扣回来
    last_out = [r for r in d2["records"] if r["kind"] == "outsource"
                and r["status"] == "已付"][0]
    fin.update(last_out["id"], {"kind": "outsource", "project_id": pid,
                                "amount": 30000, "status": "应付",
                                "date": today.isoformat(), "vendor_id": v_a})
    assert fin.data()["cards"]["net"] == 50000, "「应付」不该冲减实际收入"
    fin.update(last_out["id"], {"kind": "outsource", "project_id": pid,
                                "amount": 30000, "status": "已付",
                                "date": today.isoformat(), "vendor_id": v_a})
    assert fin.data()["cards"]["net"] == 20000

    # ---- 外包商汇总（2026-10-16 与制作人打通后新增）----
    vs = {v["vendor"]: v for v in fin.data()["vendor_summary"]}
    assert set(vs) == {"阿强", "光影工作室"}, vs
    assert vs["阿强"]["paid"] == 30000 and vs["阿强"]["pending"] == 0, vs["阿强"]
    assert vs["阿强"]["total"] == 30000 and vs["阿强"]["count"] == 1, vs["阿强"]
    assert vs["光影工作室"]["paid"] == 0 and vs["光影工作室"]["pending"] == 7000, vs
    assert vs["光影工作室"]["total"] == 7000, vs["光影工作室"]
    assert vs["光影工作室"]["projects"] == 1, vs["光影工作室"]
    assert vs["阿强"]["kind"] == "个人" and vs["光影工作室"]["kind"] == "团队", vs
    # 明细里也要带出「付给了谁」，不然展开项目看不见
    out_rows = [r for r in fin.data()["records"] if r["kind"] == "outsource"]
    assert {r["vendor"] for r in out_rows} == {"阿强", "光影工作室"}, out_rows
    assert all(r["vendor_id"] for r in out_rows), out_rows

    # ---- 必填：外包支出不给外包商 -> 拦下，理由要说清是缺了谁 ----
    try:
        fin.add({"kind": "outsource", "project_id": pid, "amount": 1,
                 "status": "已付", "date": today.isoformat()})
        raise AssertionError("外包支出没选外包商不该被放行")
    except ValueError as e:
        assert "外包商" in str(e), e
    # 别的款项类型不给外包商照常；就算硬塞一个 id 也**存不进去**（否则
    # 按人汇总时会把不属于他的到账款算进去）
    r_pay = fin.add({"kind": "payment", "project_id": pid, "amount": 123,
                     "status": "待收", "date": today.isoformat(), "vendor_id": v_a})
    pay_row = [r for r in fin.data()["records"] if r["id"] == r_pay][0]
    assert pay_row["vendor_id"] is None and pay_row["vendor"] == "", pay_row
    fin.delete(r_pay)
    # 类型从「外包支出」改成「到账」→ 人跟着清掉
    tmp_out = fin.add({"kind": "outsource", "project_id": pid, "amount": 88,
                       "status": "应付", "date": today.isoformat(), "vendor_id": v_a})
    fin.update(tmp_out, {"kind": "payment", "project_id": pid, "amount": 88,
                         "status": "待收", "date": today.isoformat()})
    sw = [r for r in fin.data()["records"] if r["id"] == tmp_out][0]
    assert sw["vendor_id"] is None, sw
    fin.delete(tmp_out)

    # ---- 外包商删除保护：三处任一挂着都不许删 ----
    for vid in (v_a, v_b):
        try:
            fin.vendor_delete(vid)
            raise AssertionError("还挂着外包支出，不该允许删除")
        except ValueError as e:
            assert "外包支出" in str(e), e
    # 类型只认两档
    try:
        fin.vendor_save(None, {"name": "野的", "kind": "公司"})
        raise AssertionError("类型只该有「个人 / 团队」两档")
    except ValueError:
        pass
    # 名字不能空
    try:
        fin.vendor_save(None, {"name": "  "})
        raise AssertionError("空名字不该被放行")
    except ValueError:
        pass
    # 重名查找：忽略大小写与首尾空白，找不到才新建
    assert fin.vendor_find_or_create("阿强") == v_a
    assert fin.vendor_find_or_create("  阿强  ") == v_a
    v_new = fin.vendor_find_or_create("新来的")
    assert v_new not in (v_a, v_b) and fin.vendor_find_or_create("新来的") == v_new
    fin.vendor_delete(v_new)                       # 没人挂 → 删得掉
    assert not [v for v in fin.vendor_options() if v["id"] == v_new]

    # ---- 环节挂载 ----
    proj = d["projects"][0]
    assert proj["node_records"] == 1, proj
    node_rows = [r for r in d["records"] if r["node_id"]]
    assert len(node_rows) == 1 and node_rows[0]["target"] == "合成", node_rows

    # ---- 状态不计入 ----
    rid = fin.add({"kind": "payment", "project_id": pid, "amount": 99999,
                   "status": "待收", "date": today.isoformat()})
    assert fin.data()["cards"]["received"] == 50000, "待收不应计入已收"
    fin.update(rid, {"kind": "payment", "project_id": pid, "amount": 99999,
                     "status": "已收", "date": today.isoformat()})
    assert fin.data()["cards"]["received"] == 149999, "改成已收后应计入"
    fin.delete(rid)
    assert fin.data()["cards"]["received"] == 50000, "删除后应回退"

    # ---- 作废合同不计入 ----
    void = fin.add({"kind": "contract", "project_id": pid, "amount": 5000,
                    "status": "作废", "date": today.isoformat()})
    assert fin.data()["cards"]["contract"] == 100000, "作废合同不该计入"
    fin.delete(void)

    # ---- 核销（2026-10-16）：只冲待收，不动合同额、不进实际收入 ----
    #     这是这轮改口径的核心：合同签了多少就是多少，收不齐由「核销」承担，
    #     所以收款率仍然是 已收 / 合同额，不会因为核销而虚高。
    c0 = fin.data()["cards"]
    assert c0["pending"] == 50000 and c0["writeoff"] == 0, c0
    wo = fin.add({"kind": "writeoff", "project_id": pid, "amount": 20000,
                  "date": today.isoformat(), "status": "已核销", "note": "尾款收不回"})
    c1 = fin.data()["cards"]
    assert c1["contract"] == 100000, "核销不该动合同额"
    assert c1["received"] == c0["received"], "核销不是到账"
    assert c1["net"] == c0["net"], "核销不该进实际收入"
    assert c1["writeoff"] == 20000, c1
    assert c1["pending"] == 30000, c1                       # 100000 − 50000 − 20000
    assert c1["rate"] == c0["rate"] == 50.0, "核销不该改收款率的分母"
    # 客户汇总、项目行都要跟着减；核销单列一列
    proj1 = [p for p in fin.data()["projects"] if p["project_id"] == pid][0]
    assert proj1["pending"] == 30000 and proj1["writeoff"] == 20000, proj1
    cli1 = [x for x in fin.data()["client_summary"] if x["client"] == "ACME 影业"][0]
    assert cli1["pending"] == 30000 and cli1["writeoff"] == 20000, cli1
    # 账龄吃的是 p["pending"]，核销过的那部分不该再挂在账上
    ag_wo = fin.data()["aging"]
    assert ag_wo["total"] == 30000, ag_wo
    # 作废的核销不计入 —— "撤销这次核销"就这一下，待收自己回来
    fin.update(wo, {"status": "作废"})
    c2 = fin.data()["cards"]
    assert c2["writeoff"] == 0 and c2["pending"] == 50000, c2
    assert fin.data()["aging"]["total"] == 50000, "撤销核销后待收要回来"
    fin.delete(wo)
    assert fin.data()["cards"]["pending"] == 50000, "删掉核销记录后待收也要回来"
    # 核销不带外包商（那不是它的字段，_vendor_id 只对 outsource 生效）
    wo2 = fin.add({"kind": "writeoff", "project_id": pid, "amount": 1,
                   "date": today.isoformat(), "vendor_id": v_a})
    assert not (fin.records(project_id=pid)[-1].get("vendor") or ""), "核销不该挂外包商"
    fin.delete(wo2)

    # ---- 账龄：合同日期 45 天前 -> 落在 31-60 档 ----
    ag = fin.data()["aging"]
    hit = [b for b in ag["buckets"] if b["label"] == "31-60 天"]
    assert hit and hit[0]["amount"] == 50000, ag
    assert ag["total"] == 50000, ag

    # ---- 客户汇总 ----
    cs = fin.data()["client_summary"]
    assert cs and cs[0]["client"] == "ACME 影业" and cs[0]["pending"] == 50000, cs

    # ---- 局部更新：只给一个字段，别的字段不能被顺手清空 ----
    # 以前 update() 一律 payload.get(k)，取不到就写空回去 —— 这里会抹掉发票号、
    # 把税率/环节清成空。前端表单每次提交都带全字段所以看不出来，只有别的调用方才踩。
    inv0 = [r for r in fin.data()["records"] if r["kind"] == "invoice"][0]
    assert inv0["invoice_no"] == "INV-001" and inv0["amount"] == 60000, inv0
    fin.update(inv0["id"], {"note": "只改备注"})
    inv1 = [r for r in fin.data()["records"] if r["id"] == inv0["id"]][0]
    assert inv1["invoice_no"] == "INV-001", inv1     # 发票号不该被抹掉
    assert inv1["amount"] == 60000, inv1             # 金额不该被清成 0
    assert inv1["kind"] == "invoice" and inv1["status"] == "已开", inv1
    assert inv1["node_id"] == inv0["node_id"], inv1  # 挂靠环节不该被摘掉
    assert inv1["note"] == "只改备注", inv1
    assert fin.data()["cards"]["invoiced"] == 60000  # 口径不受影响
    fin.update(inv0["id"], {"note": ""})

    # ---- 年份 / 币种筛选 ----
    assert fin.data(year=today.year)["cards"]["contract"] == 100000
    assert fin.data(currency="USD")["cards"]["contract"] == 0
    usd = fin.add({"kind": "contract", "project_id": pid, "amount": 2000,
                   "currency": "USD", "status": "有效", "date": today.isoformat()})
    du = fin.data(currency="USD")
    assert du["cards"]["contract"] == 2000, du["cards"]
    assert "USD" in du["currencies"] and du["other_currency"] >= 1, du
    fin.delete(usd)

    # ---- 导出 ----
    for fmt, ext in (("csv", "csv"), ("json", "json"), ("xlsx", "xlsx")):
        p = fin.export(fmt)
        assert p.endswith("." + ext), p
        assert os.path.getsize(p) > 0, p
    xlsx = [f for f in os.listdir(os.path.dirname(p)) if f.endswith(".xlsx")][-1]
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(os.path.dirname(p), xlsx))
    assert wb.sheetnames == ["款项明细", "项目汇总", "客户汇总", "外包商汇总",
                             "应收账龄"], wb.sheetnames
    assert wb["款项明细"].max_row >= 6, wb["款项明细"].max_row
    # 明细表头带了「外包商」列（在「挂到」后面）
    head1 = [c.value for c in wb["款项明细"][1]]
    assert head1 == ["日期", "类型", "状态", "客户", "项目", "挂到", "外包商",
                     "金额", "币种", "税率%", "发票号", "备注"], head1
    # 看板上不展示开票口径，但导出里要留着（数据完整），外包 / 已核销 / 实际收入
    # 几列也在。「已核销」是 2026-10-16 加的（核销只冲待收，不改合同额）
    head2 = [c.value for c in wb["项目汇总"][1]]
    assert head2 == ["项目", "客户", "合同额", "已收", "外包支出", "实际收入", "待收",
                     "已核销", "已开票", "未开票", "合同日期", "最近到账", "笔数",
                     "挂在环节"], head2
    row2 = [c.value for c in wb["项目汇总"][2]]
    assert row2[4] == 30000 and row2[5] == 20000, row2      # 外包支出 / 实际收入
    assert row2[7] == 0, row2                               # 还没结过款 → 已核销 0
    head3 = [c.value for c in wb["客户汇总"][1]]
    assert head3 == ["客户", "项目数", "合同额", "已收", "外包支出", "实际收入", "待收",
                     "已核销", "收款率%", "最早合同日"], head3
    # 外包商汇总那张表：表头 + 数据都要对
    head4 = [c.value for c in wb["外包商汇总"][1]]
    assert head4 == ["外包商", "类型", "项目数", "已付", "应付", "合计", "笔数",
                     "最近一笔"], head4
    body4 = {r[0].value: [c.value for c in r] for r in wb["外包商汇总"].iter_rows(min_row=2)}
    assert body4["阿强"][3] == 30000 and body4["阿强"][5] == 30000, body4
    assert body4["光影工作室"][4] == 7000, body4

    # ---- 客户删除保护 ----
    try:
        fin.client_delete(cid)
        raise AssertionError("客户还有项目时不该允许删除")
    except ValueError:
        pass

    # ---- 参数校验 ----
    for bad in ({"kind": "payment", "project_id": pid, "amount": 0},
                {"kind": "payment", "project_id": 0, "amount": 100},
                {"kind": "unknown", "project_id": pid, "amount": 100}):
        try:
            fin.add(bad)
            raise AssertionError(f"应当被拒绝：{bad}")
        except ValueError:
            pass

    out_dir = fin.export_dir()
    db.close()
    db2 = Db(path)
    db2.close()
    print("smoke finance ok · 导出目录:", out_dir)
    print("  cards:", money(c["contract"]), "contract /", money(c["received"]), "received")
    print("  aging:", [(b["label"], b["amount"]) for b in ag["buckets"]])
    return 0


if __name__ == "__main__":
    sys.exit(main())
